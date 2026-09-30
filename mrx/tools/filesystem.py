"""File System Agent: every mutating tool verifies its own effect on disk."""
from __future__ import annotations

import fnmatch
import mimetypes
import os
import shutil
import subprocess
import sys
import tarfile
import time
import zipfile
from pathlib import Path

from .registry import Outcome, Tool, ToolError

SPECIAL = {"desktop": "Desktop", "documents": "Documents", "downloads": "Downloads", "pictures": "Pictures",
           "music": "Music", "videos": "Videos"}
TEXT_EXT = {".txt", ".md", ".json", ".csv", ".tsv", ".log", ".ini", ".cfg", ".toml", ".yaml", ".yml", ".xml",
            ".html", ".css", ".js", ".ts", ".py", ".java", ".c", ".cpp", ".h", ".cs", ".go", ".rs", ".sh",
            ".bat", ".ps1", ".sql", ".rb", ".php", ".tsx", ".jsx"}


def special_folder(name: str) -> Path:
    home = Path.home()
    base = SPECIAL[name.lower()]
    direct = home / base
    if direct.exists():
        return direct
    onedrive = home / "OneDrive" / base  # Windows folder redirection
    if onedrive.exists():
        return onedrive
    return direct


# Locations an agent must never read or write (credentials, shell start-up files, auto-start folders, browser
# profiles, M.R.X.'s own private state). A prompt-injected agent could otherwise steal keys or plant persistence.
DENY_DIRS = {".ssh", ".aws", ".gnupg", ".kube", ".docker", ".azure", "gcloud", ".password-store", "browser-profile", "keychains"}
DENY_FILES = {".bashrc", ".bash_profile", ".profile", ".zshrc", ".zprofile", ".netrc", ".npmrc", ".pypirc", "id_rsa", "id_ed25519",
              "login data", "web data", "settings.json", "mrx.db", "authorized_keys", ".git-credentials"}
DENY_SUBSTR = ("/start menu/programs/startup", "/.config/autostart", "/library/launchagents", "/appdata/local/google/chrome/user data",
               "/appdata/local/microsoft/edge/user data", "/.mozilla/firefox")


def _denied(p: Path, rt) -> str | None:
    parts = [x.lower() for x in p.parts]
    posix = "/" + "/".join(parts[1:]) if parts else ""
    if any(sub in posix for sub in DENY_SUBSTR):
        return "auto-start / browser profile location"
    for d in DENY_DIRS:
        if d in parts[:-1] or (d == parts[-1]):
            return d
    if parts[-1] in DENY_FILES and (parts[-1] != "settings.json" or ".mrx" in parts):
        return parts[-1]
    if _within(p.resolve(), rt.settings.home.resolve()) and p.resolve() != rt.settings.home.resolve() and \
            p.resolve().relative_to(rt.settings.home.resolve()).parts[:1] not in (("trash",), ("screenshots",)):
        return "M.R.X. private data folder"
    return None


def resolve_path(rt, raw: str) -> Path:
    if not raw or not str(raw).strip():
        raise ToolError("path is empty")
    s = os.path.expandvars(os.path.expanduser(str(raw).strip().strip('"')))
    p = Path(s)
    if not p.is_absolute():
        parts = Path(s).parts
        if parts and parts[0].lower() in SPECIAL:
            p = special_folder(parts[0]).joinpath(*parts[1:])
        else:
            p = Path.home() / s
    p = Path(os.path.normpath(p))
    roots = [Path.home(), rt.settings.home] + [Path(r).expanduser() for r in rt.settings.get("filesystem.allowed_roots", [])]
    rp = p.resolve()
    if not any(_within(rp, r.resolve()) for r in roots):
        raise ToolError(f"'{p}' is outside the folders M.R.X. may touch (your home folder plus allowed roots in "
                        f"Settings → Filesystem)")
    why = _denied(p, rt)
    if why:
        raise ToolError(f"'{p.name}' is protected ({why}): credentials, start-up files and M.R.X. private data are off-limits to the agent")
    return p


def _within(p: Path, root: Path) -> bool:
    try:
        p.relative_to(root)
        return True
    except ValueError:
        return False


def _protected(rt, p: Path) -> None:
    rp = p.resolve()
    roots = [Path.home().resolve(), rt.settings.home.resolve(), Path(rp.anchor)] + \
            [Path(r).expanduser().resolve() for r in rt.settings.get("filesystem.allowed_roots", [])]
    roots += [special_folder(n).resolve() for n in SPECIAL]
    if rp in roots:
        raise ToolError(f"refusing to modify protected folder '{rp}'")


def _size(p: Path) -> int:
    if p.is_file():
        return p.stat().st_size
    return sum(f.stat().st_size for f in p.rglob("*") if f.is_file())


def create_folder(rt, path: str):
    p = resolve_path(rt, path)
    existed = p.is_dir()
    p.mkdir(parents=True, exist_ok=True)
    ok = p.is_dir()
    return Outcome({"path": str(p), "already_existed": existed}, ok,
                   f"folder exists on disk: {ok}")


def create_file(rt, path: str, content: str = "", overwrite: bool = False):
    p = resolve_path(rt, path)
    if p.exists() and not overwrite:
        raise ToolError(f"'{p}' already exists (pass overwrite=true to replace it)")
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(content, encoding="utf-8")
    ok = p.is_file() and p.read_text(encoding="utf-8") == content
    return Outcome({"path": str(p), "bytes": p.stat().st_size if p.exists() else 0}, ok,
                   "file re-read from disk and content matches" if ok else "content on disk differs")


def read_file(rt, path: str, max_chars: int = 20000):
    p = resolve_path(rt, path)
    if not p.is_file():
        raise ToolError(f"'{p}' is not a file")
    ext = p.suffix.lower()
    text: str
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ToolError("PDF reading needs the 'pypdf' package (pip install pypdf)")
        text = "\n".join((pg.extract_text() or "") for pg in PdfReader(str(p)).pages)
    elif ext == ".docx":
        try:
            import docx
        except ImportError:
            raise ToolError("DOCX reading needs 'python-docx' (pip install python-docx)")
        text = "\n".join(par.text for par in docx.Document(str(p)).paragraphs)
    elif ext == ".xlsx":
        try:
            import openpyxl
        except ImportError:
            raise ToolError("XLSX reading needs 'openpyxl' (pip install openpyxl)")
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"## {ws.title}")
            for row in ws.iter_rows(values_only=True):
                lines.append("\t".join("" if c is None else str(c) for c in row))
        text = "\n".join(lines)
    elif ext in TEXT_EXT or (mimetypes.guess_type(str(p))[0] or "").startswith("text"):
        text = p.read_text(encoding="utf-8", errors="replace")
    else:
        raise ToolError(f"cannot read '{ext or 'unknown'}' files as text (type: "
                        f"{mimetypes.guess_type(str(p))[0] or 'unknown'}); use file_info or open_path")
    truncated = len(text) > max_chars
    return {"path": str(p), "content": text[:max_chars], "chars": len(text), "truncated": truncated}


def modify_file(rt, path: str, content: str | None = None, find: str | None = None, replace: str | None = None):
    p = resolve_path(rt, path)
    if not p.is_file():
        raise ToolError(f"'{p}' does not exist; use create_file")
    if content is not None:
        new = content
    elif find is not None and replace is not None:
        old = p.read_text(encoding="utf-8")
        if find not in old:
            raise ToolError("text to replace was not found in the file")
        new = old.replace(find, replace)
    else:
        raise ToolError("provide 'content' (full replacement) or both 'find' and 'replace'")
    p.write_text(new, encoding="utf-8")
    ok = p.read_text(encoding="utf-8") == new
    return Outcome({"path": str(p), "bytes": len(new.encode())}, ok, "re-read from disk matches: %s" % ok)


def append_file(rt, path: str, content: str):
    p = resolve_path(rt, path)
    before = p.stat().st_size if p.exists() else 0
    p.parent.mkdir(parents=True, exist_ok=True)
    with p.open("a", encoding="utf-8") as f:
        f.write(content)
    after = p.stat().st_size
    ok = after == before + len(content.encode())
    return Outcome({"path": str(p), "bytes": after}, ok, f"size grew {before}→{after}")


def copy_file(rt, src: str, dst: str):
    s, d = resolve_path(rt, src), resolve_path(rt, dst)
    if not s.exists():
        raise ToolError(f"source '{s}' does not exist")
    if d.is_dir():
        d = d / s.name
    if d.exists():
        raise ToolError(f"destination '{d}' already exists")
    d.parent.mkdir(parents=True, exist_ok=True)
    if s.is_dir():
        shutil.copytree(s, d)
    else:
        shutil.copy2(s, d)
    ok = d.exists() and _size(d) == _size(s)
    return Outcome({"path": str(d), "src": str(s)}, ok, f"copy exists, size {_size(d) if d.exists() else 0} vs {_size(s)}")


def move_file(rt, src: str, dst: str):
    s, d = resolve_path(rt, src), resolve_path(rt, dst)
    if not s.exists():
        raise ToolError(f"source '{s}' does not exist")
    _protected(rt, s)
    if d.is_dir():
        d = d / s.name
    if d.exists():
        raise ToolError(f"destination '{d}' already exists")
    d.parent.mkdir(parents=True, exist_ok=True)
    size = _size(s)
    shutil.move(str(s), str(d))
    ok = d.exists() and not s.exists() and _size(d) == size
    return Outcome({"path": str(d), "src": str(s)}, ok, f"destination exists: {d.exists()}, source gone: {not s.exists()}")


def rename_file(rt, path: str, new_name: str):
    p = resolve_path(rt, path)
    if not p.exists():
        raise ToolError(f"'{p}' does not exist")
    if any(c in new_name for c in "/\\"):
        raise ToolError("new_name must be a bare name, not a path (use move_file to relocate)")
    _protected(rt, p)
    d = p.with_name(new_name)
    if d.exists():
        raise ToolError(f"'{d}' already exists")
    p.rename(d)
    ok = d.exists() and not p.exists()
    return Outcome({"path": str(d), "src": str(p)}, ok, f"new name exists: {d.exists()}, old name gone: {not p.exists()}")


def _remove(rt, p: Path, permanent: bool) -> str:
    _protected(rt, p)
    use_trash = rt.settings.get("filesystem.trash_instead_of_delete", True) and not permanent
    if use_trash:
        dest_dir = rt.settings.home / "trash" / time.strftime("%Y%m%d-%H%M%S")
        dest_dir.mkdir(parents=True, exist_ok=True)
        dest = dest_dir / p.name
        shutil.move(str(p), str(dest))
        return f"moved to trash: {dest}"
    shutil.rmtree(p) if p.is_dir() else p.unlink()
    return "permanently deleted"


def delete_file(rt, path: str, permanent: bool = False):
    p = resolve_path(rt, path)
    if not p.is_file():
        raise ToolError(f"'{p}' is not a file")
    how = _remove(rt, p, permanent)
    ok = not p.exists()
    return Outcome({"path": str(p), "how": how}, ok, f"no longer exists on disk: {ok}")


def delete_folder(rt, path: str, permanent: bool = False):
    p = resolve_path(rt, path)
    if not p.is_dir():
        raise ToolError(f"'{p}' is not a folder")
    count = sum(1 for _ in p.rglob("*"))
    how = _remove(rt, p, permanent)
    ok = not p.exists()
    return Outcome({"path": str(p), "items_removed": count, "how": how}, ok, f"no longer exists on disk: {ok}")


def _entry(e: Path) -> dict:
    try:
        st = e.stat()
        return {"name": e.name, "path": str(e), "is_dir": e.is_dir(), "size": st.st_size if e.is_file() else None,
                "modified": st.st_mtime}
    except OSError:
        return {"name": e.name, "path": str(e), "is_dir": False, "size": None, "modified": None}


def list_directory(rt, path: str = "~", include_hidden: bool = False, limit: int = 500):
    p = resolve_path(rt, path)
    if not p.is_dir():
        raise ToolError(f"'{p}' is not a folder")
    items = [e for e in sorted(p.iterdir(), key=lambda x: (not x.is_dir(), x.name.lower()))
             if include_hidden or not e.name.startswith(".")]
    return {"path": str(p), "parent": str(p.parent), "count": len(items), "truncated": len(items) > limit,
            "entries": [_entry(e) for e in items[:limit]]}


def search_files(rt, query: str, root: str = "~", content: bool = False, limit: int = 50, max_scanned: int = 200000):
    base = resolve_path(rt, root)
    if not base.is_dir():
        raise ToolError(f"'{base}' is not a folder")
    q = query.lower()
    pattern = query if any(c in query for c in "*?[") else None
    hits, scanned = [], 0
    skip = {"node_modules", ".git", "__pycache__", "AppData", ".cache", "$RECYCLE.BIN"}
    for dp, dns, fns in os.walk(base):
        dns[:] = [d for d in dns if d not in skip]
        for n in dns + fns:
            scanned += 1
            match = fnmatch.fnmatch(n.lower(), pattern.lower()) if pattern else q in n.lower()
            fp = Path(dp) / n
            if not match and content and n in fns and Path(n).suffix.lower() in TEXT_EXT:
                try:
                    if fp.stat().st_size < 2_000_000 and q in fp.read_text("utf-8", errors="ignore").lower():
                        match = True
                except OSError:
                    pass
            if match:
                hits.append(_entry(fp))
                if len(hits) >= limit:
                    return {"query": query, "root": str(base), "matches": hits, "truncated": True, "scanned": scanned}
            if scanned >= max_scanned:
                return {"query": query, "root": str(base), "matches": hits, "truncated": True, "scanned": scanned}
    return {"query": query, "root": str(base), "matches": hits, "truncated": False, "scanned": scanned}


def file_info(rt, path: str):
    p = resolve_path(rt, path)
    if not p.exists():
        raise ToolError(f"'{p}' does not exist")
    st = p.stat()
    return {"path": str(p), "exists": True, "is_dir": p.is_dir(), "size": _size(p),
            "mime": mimetypes.guess_type(str(p))[0], "extension": p.suffix.lower(),
            "modified": st.st_mtime, "created": st.st_ctime}


def open_path(rt, path: str, reveal: bool = False):
    p = resolve_path(rt, path)
    if not p.exists():
        raise ToolError(f"'{p}' does not exist")
    target = p.parent if (reveal and p.is_file()) else p
    if sys.platform.startswith("win"):
        if reveal and p.is_file():
            subprocess.Popen(["explorer", f"/select,{p}"])
        else:
            os.startfile(str(target))  # type: ignore[attr-defined]
    elif sys.platform == "darwin":
        subprocess.Popen(["open", "-R", str(p)] if reveal else ["open", str(target)])
    else:
        opener = shutil.which("xdg-open")
        if not opener:
            raise ToolError("no 'xdg-open' available on this system to open files")
        subprocess.Popen([opener, str(target)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    # The OS accepted the open request; whether a window appeared cannot be observed generically.
    return Outcome({"path": str(target)}, None, "OS accepted the open request; window not independently confirmed")


def archive_files(rt, paths: list, dest: str):
    d = resolve_path(rt, dest)
    if d.exists():
        raise ToolError(f"'{d}' already exists")
    srcs = [resolve_path(rt, x) for x in paths]
    for s in srcs:
        if not s.exists():
            raise ToolError(f"'{s}' does not exist")
    d.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(d, "w", zipfile.ZIP_DEFLATED) as z:
        for s in srcs:
            if s.is_dir():
                for f in s.rglob("*"):
                    if f.is_file():
                        z.write(f, Path(s.name) / f.relative_to(s))
            else:
                z.write(s, s.name)
    ok = zipfile.is_zipfile(d)
    with zipfile.ZipFile(d) as z:
        n = len(z.namelist())
    return Outcome({"path": str(d), "entries": n}, ok and n > 0, f"valid zip with {n} entries")


def extract_archive(rt, archive: str, dest: str):
    a, d = resolve_path(rt, archive), resolve_path(rt, dest)
    if not a.is_file():
        raise ToolError(f"'{a}' is not a file")
    d.mkdir(parents=True, exist_ok=True)
    root = d.resolve()
    if zipfile.is_zipfile(a):
        with zipfile.ZipFile(a) as z:
            for m in z.namelist():
                if not _within((d / m).resolve(), root):
                    raise ToolError(f"unsafe path in archive: {m}")
            z.extractall(d)
            names = z.namelist()
    elif tarfile.is_tarfile(a):
        with tarfile.open(a) as t:
            names = t.getnames()
            for m in names:
                if not _within((d / m).resolve(), root):
                    raise ToolError(f"unsafe path in archive: {m}")
            t.extractall(d, filter="data")
    else:
        raise ToolError("unsupported archive type (zip and tar are supported)")
    missing = [n for n in names if not (d / n).exists()]
    return Outcome({"path": str(d), "entries": len(names)}, not missing,
                   f"{len(names) - len(missing)}/{len(names)} entries present on disk")


def _p(**props):
    return props


def tools() -> list[Tool]:
    S = {"type": "string"}
    P = {"type": "string", "description": "File/folder path. Relative paths start in the home folder; "
                                          "'Desktop/…', 'Documents/…', 'Downloads/…' are understood."}
    F = "plugin:filesystem"
    return [
        Tool("create_folder", "Create a folder (and parents). Verified on disk.", {"path": P}, ["path"], create_folder, scope=F),
        Tool("create_file", "Create a text file with content. Verified by re-reading.",
             {"path": P, "content": S, "overwrite": {"type": "boolean"}}, ["path"], create_file, scope=F),
        Tool("read_file", "Read a file (text, code, JSON, CSV, PDF, DOCX, XLSX).",
             {"path": P, "max_chars": {"type": "integer"}}, ["path"], read_file, scope=F),
        Tool("modify_file", "Replace a file's content, or find/replace text inside it.",
             {"path": P, "content": S, "find": S, "replace": S}, ["path"], modify_file, scope=F),
        Tool("append_file", "Append text to a file.", {"path": P, "content": S}, ["path", "content"], append_file, scope=F),
        Tool("copy_file", "Copy a file or folder.", {"src": P, "dst": P}, ["src", "dst"], copy_file, scope=F),
        Tool("move_file", "Move a file or folder.", {"src": P, "dst": P}, ["src", "dst"], move_file, risk=2, scope=F),
        Tool("rename_file", "Rename a file or folder in place.", {"path": P, "new_name": S}, ["path", "new_name"],
             rename_file, risk=2, scope=F),
        Tool("delete_file", "Delete a file (moved to the M.R.X. trash so it is recoverable).",
             {"path": P, "permanent": {"type": "boolean"}}, ["path"], delete_file, risk=2, scope=F),
        Tool("delete_folder", "Delete a folder and everything in it (moved to the M.R.X. trash).",
             {"path": P, "permanent": {"type": "boolean"}}, ["path"], delete_folder, risk=3, scope=F),
        Tool("list_directory", "List a folder's contents.", {"path": P, "include_hidden": {"type": "boolean"},
                                                            "limit": {"type": "integer"}}, [], list_directory, scope=F),
        Tool("search_files", "Search files by name (or glob) and optionally by text content.",
             {"query": S, "root": P, "content": {"type": "boolean"}, "limit": {"type": "integer"}},
             ["query"], search_files, scope=F),
        Tool("file_info", "Get metadata and detected type of a file/folder.", {"path": P}, ["path"], file_info, scope=F),
        Tool("open_path", "Open a file/folder with the default app, or reveal it in the file manager.",
             {"path": P, "reveal": {"type": "boolean"}}, ["path"], open_path, scope=F),
        Tool("archive_files", "Compress files/folders into a zip.",
             {"paths": {"type": "array", "items": S}, "dest": P}, ["paths", "dest"], archive_files, scope=F),
        Tool("extract_archive", "Extract a zip/tar archive (path-traversal safe).",
             {"archive": P, "dest": P}, ["archive", "dest"], extract_archive, scope=F),
    ]
